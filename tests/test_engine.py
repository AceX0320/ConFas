"""Cross-platform unit tests for the ConFas engine (run anywhere)."""
import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))
os.environ["CONFAS_HOME"] = str(Path(__file__).parent / "_tmp_home")

import pytest  # noqa: E402
from confas.engine import ConversionEngine, newest_target_for, _routes_for  # noqa: E402


@pytest.fixture()
def engine():
    return ConversionEngine(enable_ml=False)


def test_newest_generation_targets():
    assert newest_target_for("pdf")[0] == "docx"      # latest Word, not .doc
    assert newest_target_for("doc")[0] == "docx"      # legacy -> modern
    assert newest_target_for("xls")[0] == "xlsx"
    assert newest_target_for("ppt")[0] == "pptx"
    assert "doc" not in newest_target_for("docx")     # never downgrade


def test_routes_prefer_fastest(engine):
    routes = [r.name for r in _routes_for(".png", ".jpg")]
    assert routes and routes[0] == "pillow"           # C-speed path first


def test_png_to_jpg_roundtrip(tmp_path, engine):
    from PIL import Image

    src = tmp_path / "a.png"
    Image.new("RGBA", (64, 32), (10, 200, 30, 255)).save(src)
    res = engine.convert(src, "jpg")
    assert res.ok, res.error
    out = tmp_path / "a.jpg"
    assert out.exists()
    with Image.open(out) as im:
        assert im.size == (64, 32)


def test_txt_to_docx(tmp_path, engine):
    src = tmp_path / "note.txt"
    src.write_text("# Title\nHello **world**\n- one\n- two\n", encoding="utf-8")
    res = engine.convert(src, "docx")
    assert res.ok, res.error
    import docx

    d = docx.Document(str(tmp_path / "note.docx"))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Title" in text and "one" in text
    # verify it is a real OOXML zip (Word 2007+ generation)
    import zipfile

    assert zipfile.is_zipfile(tmp_path / "note.docx")


def test_docx_to_txt(tmp_path, engine):
    import docx

    d = docx.Document()
    d.add_paragraph("Alpha")
    d.add_paragraph("Beta")
    f = tmp_path / "x.docx"
    d.save(f)
    res = engine.convert(f, "txt")
    assert res.ok, res.error
    assert "Alpha" in (tmp_path / "x.txt").read_text(encoding="utf-8-sig")


def test_missing_source(engine):
    res = engine.convert("/nope/missing.pdf", "docx")
    assert not res.ok


def test_unsupported_pair_returns_error(tmp_path, engine):
    f = tmp_path / "weird.zzz"
    f.write_bytes(b"\x00\x01")
    res = engine.convert(f, "qqq")
    assert not res.ok  # must fail loudly rather than produce garbage
